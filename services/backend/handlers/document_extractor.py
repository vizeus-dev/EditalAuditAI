#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
Document Extractor (Offline-First) & Heurísticas de Domínio Editalício
Extrai texto limpo de PDF, DOCX e TXT sem chamadas de rede externas e analisa metadados de certame.
Padrão Ponytail (Stdlib / pypdf / python-docx) + Padrão Matt Pocock (Tipagem Estrutural).
"""

import io
import re
from typing import Dict, Any, Optional

try:
    import pypdf
    PYPDF_AVAILABLE = True
except ImportError:
    PYPDF_AVAILABLE = False

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


def detect_edital_metadata(text: str) -> Dict[str, str]:
    """
    Identifica por heurísticas determinísticas o órgão promotor, título do edital,
    teto orçamentário e prazos de submissão a partir do texto do edital.
    """
    metadata: Dict[str, str] = {
        "institution": "",
        "title": "",
        "budget": "",
        "deadlines": "",
        "year": "2026"
    }
    
    if not text:
        return metadata

    # 1. Órgão / Instituição
    inst_match = re.search(
        r'(?:SECRETARIA|MINISTÉRIO|FUNDAÇÃO|PREFEITURA|GOVERNO)\s+(?:MUNICIPAL|ESTADUAL|NACIONAL|DA|DO|DE)?\s+([A-ZÀ-Ú\s]{3,50})',
        text,
        re.IGNORECASE
    )
    if inst_match:
        metadata["institution"] = inst_match.group(0).strip()
    else:
        metadata["institution"] = "Órgão Promotor do Edital"

    # 2. Título do Edital
    title_match = re.search(
        r'(?:EDITAL(?:\s+DE)?\s+[^\n\r]{5,70}|CHAMAMENTO\s+PÚBLICO[^\n\r]{5,70}|FESTIVAL\s+[^\n\r]{5,70}|PROAC[^\n\r]{3,40})',
        text,
        re.IGNORECASE
    )
    if title_match:
        clean_title = re.sub(r'[\r\n]+', ' ', title_match.group(0)).strip()
        metadata["title"] = clean_title
    else:
        metadata["title"] = "Projeto Cultural de Fomento"

    # 3. Objeto do Edital
    objeto_match = re.search(
        r'(?:objeto|finalidade|finalidades)[\s\:\-]+([^\n\r]{10,200})',
        text,
        re.IGNORECASE
    )
    if objeto_match:
        metadata["objeto"] = objeto_match.group(1).strip()

    # 4. Orçamento / Teto Estimado (Análise contextual ponderada para filtrar taxas e capturar o teto real)
    budget_candidates = []
    for match in re.finditer(r'(R\$\s*[\d\.,]+)', text):
        raw_val = match.group(1).rstrip('.')
        start_ctx = max(0, match.start() - 100)
        ctx_before = text[start_ctx:match.start()].lower()

        # Converte para número para descartar valores zerados (ex: taxa R$ 0,00)
        num_clean = re.sub(r'[^\d]', '', raw_val)
        try:
            val_float = float(num_clean) / (100.0 if ',' in raw_val else 1.0)
        except ValueError:
            val_float = 0.0

        if val_float <= 0:
            continue

        # Pontuação de relevância de contexto
        weight = 0
        if any(k in ctx_before for k in ["teto", "máximo", "maximo", "limite", "solicitar", "projeto"]):
            weight += 10
        if any(k in ctx_before for k in ["valor total", "recursos", "orçamento", "orcamento", "global", "disponibilizado"]):
            weight += 5
        if "taxa" in ctx_before or "inscrição" in ctx_before or "inscricao" in ctx_before:
            weight -= 10

        budget_candidates.append((weight, val_float, raw_val))

    if budget_candidates:
        budget_candidates.sort(key=lambda x: (x[0], x[1]), reverse=True)
        metadata["budget"] = budget_candidates[0][2].strip()
    else:
        metadata["budget"] = ""

    # 5. Prazos / Inscrições
    deadline_match = re.search(
        r'(?:inscriç(?:ões|ao|ão)|prazo|envio|submissão).*?\b(\d{1,2}[\/\.-]\d{1,2}[\/\.-]\d{2,4})\b',
        text,
        re.IGNORECASE
    )
    if deadline_match:
        metadata["deadlines"] = deadline_match.group(1).strip()

    # 6. Ano
    year_match = re.search(r'\b(202[4-9]|203[0-5])\b', text)
    if year_match:
        metadata["year"] = year_match.group(1)

    return metadata


def generate_proposal_draft_suggestion(edital_text: str, notes_text: str = "") -> Dict[str, str]:
    """
    Gera um esboço inicial inteligente para os campos da proposta ABNT
    cruzando os requisitos identificados no edital com as notas fornecidas pelo proponente.
    Retorna tanto as chaves canônicas quanto os aliases esperados pelo frontend e compilador PDF.
    """
    meta = detect_edital_metadata(edital_text)
    notes_clean = notes_text.strip()
    
    inst = meta.get("institution") or "o órgão de fomento"
    title = meta.get("title") or "o edital vigente"
    objeto = meta.get("objeto") or ""

    apresentacao = (
        f"A presente proposta técnica tem por finalidade a execução do projeto '{title}', "
        f"submetido ao processo seletivo promovido por {inst}."
    )
    if objeto:
        apresentacao += f" O certame destina-se a {objeto}."
    if notes_clean:
        apresentacao += f" Premissa central do proponente: {notes_clean}."

    justificativa = (
        f"O presente projeto cultural alinha-se diretamente às diretrizes e objetivos estabelecidos por {inst} "
        f"no âmbito do certame {title}. A iniciativa busca democratizar o acesso aos bens culturais, "
        f"fortalecer a cadeia produtiva local e valorizar a diversidade de expressões artísticas territoriais."
    )
    if notes_clean:
        justificativa += f" Especificamente, o projeto enfatiza as seguintes premissas do proponente: {notes_clean}."

    objetivos = (
        "Objetivo Geral: Realizar a produção, circulação e difusão das atividades artístico-culturais previstas, "
        "promovendo o engajamento comunitário e a fruição estética de qualidade.\n\n"
        "Metas Específicas:\n"
        "1. Executar 100% das etapas operacionais de pré-produção, execução e prestação de contas com rigor documental.\n"
        "2. Garantir atendimento prioritário a públicos com menor acesso histórico a bens culturais.\n"
        "3. Produzir registro audiovisual em alta definição para fins de memória, acervo público e difusão gratuita."
    )

    metodologia = (
        "O plano de trabalho está estruturado em 3 etapas operacionais integradas:\n"
        "Etapa 1 - Pré-Produção: Articulação de parcerias locais, contratação formal da equipe técnica e ensaios gerais.\n"
        "Etapa 2 - Produção e Execução: Montagem de rider técnico, realização das ações artísticas e acolhimento do público.\n"
        "Etapa 3 - Pós-Produção: Relatório de cumprimento de metas, consolidação de notas fiscais e prestação de contas."
    )

    cronograma = (
        "Mês 1 a 2: Alinhamento de pré-produção, reservas de espaço e contratação de serviços especializados.\n"
        "Mês 3: Divulgação institucional, campanha de comunicação e início dos ensaios.\n"
        "Mês 4: Execução das apresentações e oficinas presenciais.\n"
        "Mês 5: Sistematização de dados de monitoramento, laudos e finalização da prestação de contas."
    )

    acessibilidade = (
        "Conforme preconizado pela Lei Brasileira de Inclusão (Lei 13.146/2015) e a ABNT NBR 9050:\n"
        "- Medidas Comunicacionais: Todas as apresentações contarão com Intérprete de Libras e audiodescrição gravada/ao vivo.\n"
        "- Medidas Físicas: Espaços com rampas de acesso, sanitários adaptados, piso tátil e reserva de assentos para PCD.\n"
        "- Medidas Atitudinais: Equipe capacitada para recepção inclusiva e acolhimento prioritário."
    )

    publico = (
        "Público Direto: Estudantes da rede pública de ensino, famílias e trabalhadores da comunidade local.\n"
        "Público Indireto: Sociedade em geral alcançada pelas divulgações e registros audiovisuais abertos na internet.\n"
        "Estimativa de Alcance: Mínimo de 500 participantes presenciais e 2.000 visualizações online."
    )

    contrapartida = (
        "Como compromisso de retorno social (Art. 14 da Lei 14.903/2024):\n"
        "1. Realização de oficinas formativas gratuitas voltadas para jovens e educadores da região.\n"
        "2. Doação de exemplares/registros documentais para a biblioteca pública municipal.\n"
        "3. Entrada 100% gratuita para todas as atividades propostas."
    )
    if notes_clean:
        contrapartida += f"\n4. Ações específicas do proponente: {notes_clean}"

    comunicacao = (
        "Plano integrado de comunicação com foco em mídias digitais e assessoria de imprensa comunitária. "
        "Aplicação obrigatória das marcas e créditos institucionais de fomento em todos os materiais de divulgação."
    )

    ficha_tecnica = (
        "Coordenação Geral e Artística: Profissional com mais de 5 anos de atuação comprovada no setor.\n"
        "Produção Executiva: Responsável por licenças, contratos e logística.\n"
        "Equipe Técnica: Operadores de áudio/luz, cenotécnico e equipe de acessibilidade (Libras e Audiodescrição)."
    )

    monitoramento = (
        "Aferição por meio de lista de presença com registro de autodeclaração, questionários de satisfação (escala Likert) "
        "e relatório fotográfico/videográfico geolocalizado de cada evento."
    )

    compliance = (
        "Regularidade fiscal atestada por CND, CNDT, FGTS e ausência de restrições no Cadin/CEIS. "
        "Respeito integral aos direitos autorais (recolhimento de Ecad quando aplicável) e normas sanitárias locais."
    )

    sustentabilidade = (
        "Implementação do Programa ESG Local: Coleta seletiva de resíduos gerados durante as ações, eliminação de copos "
        "plásticos descartáveis e compensação de carbono através do plantio de mudas nativas em área urbana."
    )

    rider = (
        "Sistema de sonorização PA estéreo balanceado, microfonação sem fio homologada pela Anatel, iluminação cênica em LED "
        "de baixo consumo energético e extintores de incêndio inspecionados no local das atividades."
    )

    return {
        "apresentacao": apresentacao,
        "justificativa": justificativa,
        "objetivos": objetivos,
        "metodologia": metodologia,
        "cronograma": cronograma,
        "acessibilidade": acessibilidade,
        "publico": publico,
        "democratizacao": publico,
        "contrapartida": contrapartida,
        "ficha_tecnica": ficha_tecnica,
        "equipe": ficha_tecnica,
        "comunicacao": comunicacao,
        "monitoramento": monitoramento,
        "compliance": compliance,
        "sustentabilidade": sustentabilidade,
        "rider": rider
    }


def extract_text_and_metadata(file_bytes: bytes, filename: str, max_pages: int = 300) -> Dict[str, Any]:
    """
    Processa arquivos PDF, DOCX ou TXT diretamente em memória.
    Retorna o texto extraído, contagem de palavras, páginas e metadados de domínio.
    """
    fn_lower = filename.lower()
    full_text = ""
    pages_count = 1
    file_format = "txt"
    error_msg: Optional[str] = None

    try:
        if fn_lower.endswith(".pdf"):
            file_format = "pdf"
            if not PYPDF_AVAILABLE:
                return {"error": "Biblioteca pypdf não disponível no ambiente.", "text": ""}
            
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages_count = len(reader.pages)
            extracted_pages = []
            
            # Limite de segurança de páginas
            limit = min(pages_count, max_pages)
            for i in range(limit):
                page_text = reader.pages[i].extract_text()
                if page_text:
                    extracted_pages.append(page_text.strip())
            
            full_text = "\n\n".join(extracted_pages)

        elif fn_lower.endswith(".docx"):
            file_format = "docx"
            if not DOCX_AVAILABLE:
                return {"error": "Biblioteca python-docx não disponível no ambiente.", "text": ""}
            
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Extrai também tabelas
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        paragraphs.append(" | ".join(row_cells))
            
            full_text = "\n\n".join(paragraphs)
            pages_count = max(1, len(paragraphs) // 5)

        else:
            file_format = "txt"
            try:
                full_text = file_bytes.decode('utf-8')
            except UnicodeDecodeError:
                full_text = file_bytes.decode('latin-1', errors='replace')
            pages_count = 1

    except Exception as e:
        error_msg = f"Falha na leitura do arquivo {filename}: {str(e)}"
        return {
            "error": error_msg,
            "text": "",
            "filename": filename,
            "format": file_format,
            "pages_count": 0,
            "words_count": 0
        }

    clean_text = full_text.strip()
    metadata = detect_edital_metadata(clean_text)
    
    return {
        "text": clean_text,
        "filename": filename,
        "format": file_format,
        "pages_count": pages_count,
        "words_count": len(clean_text.split()),
        "chars_count": len(clean_text),
        "metadata": metadata
    }
