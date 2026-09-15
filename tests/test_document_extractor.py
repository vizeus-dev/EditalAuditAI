#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
Testes Unitários: Document Extractor (Offline-First)
Valida extração de texto de PDF, DOCX, TXT, limites de segurança e heurísticas de metadados.
Padrão Ponytail (Stdlib / pypdf / python-docx) + Matt Pocock (Tipagem estrita).
"""

import io
import unittest
import pypdf
import docx

from services.backend.handlers.document_extractor import (
    extract_text_and_metadata,
    detect_edital_metadata,
    generate_proposal_draft_suggestion
)


class TestDocumentExtractor(unittest.TestCase):

    def test_extract_txt_utf8(self):
        sample_text = (
            "SECRETARIA MUNICIPAL DE CULTURA DE BELO HORIZONTE\n"
            "EDITAL DE CHAMAMENTO PÚBLICO Nº 05/2026\n"
            "FESTIVAL SONS DAS GERAIS\n"
            "Valor Total Disponibilizado: R$ 250.000,00\n"
            "Inscrições abertas até 30/11/2026."
        )
        file_bytes = sample_text.encode('utf-8')
        result = extract_text_and_metadata(file_bytes, "edital_sons.txt")

        self.assertIn("SECRETARIA MUNICIPAL DE CULTURA", result["text"])
        self.assertEqual(result["format"], "txt")
        self.assertEqual(result["pages_count"], 1)
        self.assertGreater(result["words_count"], 10)
        self.assertIn("metadata", result)
        self.assertIn("250.000", result["metadata"]["budget"])

    def test_extract_txt_latin1_fallback(self):
        sample_text = "FUNDAÇÃO CULTURAL DE NITERÓI - EDITAL ARTES VISUAIS - ORÇAMENTO R$ 80.000,00"
        file_bytes = sample_text.encode('latin-1')
        result = extract_text_and_metadata(file_bytes, "edital_niteroi.txt")

        self.assertIn("FUNDAÇÃO CULTURAL", result["text"])
        self.assertEqual(result["format"], "txt")

    def test_extract_docx_paragraphs_and_tables(self):
        doc = docx.Document()
        doc.add_heading("MINISTÉRIO DA CULTURA", level=1)
        doc.add_paragraph("EDITAL NACIONAL DE TEATRO E CIRCO 2026")
        doc.add_paragraph("O valor máximo por projeto será de R$ 150.000,00.")
        
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Categoria"
        table.rows[0].cells[1].text = "Valor"
        table.rows[1].cells[0].text = "Circulação Regional"
        table.rows[1].cells[1].text = "R$ 150.000,00"

        doc_io = io.BytesIO()
        doc.save(doc_io)
        doc_bytes = doc_io.getvalue()

        result = extract_text_and_metadata(doc_bytes, "edital_teatro.docx")
        self.assertEqual(result["format"], "docx")
        self.assertIn("MINISTÉRIO DA CULTURA", result["text"])
        self.assertIn("Circulação Regional", result["text"])
        self.assertIn("R$ 150.000,00", result["metadata"]["budget"])

    def test_extract_pdf_mock_or_stream(self):
        writer = pypdf.PdfWriter()
        page = writer.add_blank_page(width=595, height=842)
        pdf_io = io.BytesIO()
        writer.write(pdf_io)
        pdf_bytes = pdf_io.getvalue()

        result = extract_text_and_metadata(pdf_bytes, "documento_em_branco.pdf")
        self.assertEqual(result["format"], "pdf")
        self.assertEqual(result["pages_count"], 1)
        self.assertIsInstance(result["text"], str)

    def test_detect_edital_metadata_heuristics(self):
        text = (
            "GOVERNO DO ESTADO DE SÃO PAULO\n"
            "SECRETARIA DA CULTURA, ECONOMIA E INDÚSTRIA CRIATIVAS\n"
            "EDITAL PROAC Nº 22/2026 - CIRCULAÇÃO DE ESPETÁCULOS\n"
            "OBJETO: Seleção de projetos de artes cênicas para circulação no interior paulista.\n"
            "Taxa de inscrição: R$ 0,00.\n"
            "O teto orçamentário do projeto não poderá exceder R$ 120.000,00.\n"
            "As inscrições serão recebidas até 15/10/2026."
        )
        meta = detect_edital_metadata(text)
        self.assertIn("SECRETARIA", meta["institution"].upper())
        self.assertIn("PROAC", meta["title"].upper())
        self.assertIn("120.000", meta["budget"])
        self.assertIn("15/10/2026", meta["deadlines"])
        self.assertIn("artes cênicas", meta.get("objeto", "").lower())

    def test_generate_proposal_draft_suggestion(self):
        text = "SECRETARIA DE CULTURA - EDITAL FOMENTO CULTURAL - TEATRO DE RUA - OBJETO: Difusão cultural em praças - VALOR R$ 90.000,00"
        notes = "Foco em comunidades quilombolas e oficinas para jovens."
        draft = generate_proposal_draft_suggestion(text, notes)

        # Seção 1 (Apresentação)
        self.assertIn("apresentacao", draft)
        self.assertIn("teatro de rua", draft["apresentacao"].lower())

        # Seções e aliases
        self.assertIn("justificativa", draft)
        self.assertIn("objetivos", draft)
        self.assertIn("acessibilidade", draft)
        self.assertIn("contrapartida", draft)
        self.assertIn("publico", draft)
        self.assertIn("democratizacao", draft)
        self.assertEqual(draft["publico"], draft["democratizacao"])
        self.assertIn("ficha_tecnica", draft)
        self.assertIn("equipe", draft)
        self.assertEqual(draft["ficha_tecnica"], draft["equipe"])
        self.assertIn("rider", draft)

        self.assertIn("quilombolas", draft["justificativa"].lower())
        self.assertIn("oficinas", draft["contrapartida"].lower())

    def test_corrupted_data_safety(self):
        corrupted = b"NOT_A_REAL_PDF_JUST_TRASH_123456"
        result = extract_text_and_metadata(corrupted, "corrupto.pdf")
        self.assertIn("error", result)
        self.assertEqual(result["text"], "")


if __name__ == '__main__':
    unittest.main()
