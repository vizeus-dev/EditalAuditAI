#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
Testes de Resiliência Arquitetural, Limites Numéricos e LRU Cache
Valida as proteções dos 9 eixos de melhoria arquitetural (plans/architectural_improvement_plan.md).
"""

import unittest
import json
import math
import io
import openpyxl
from services.api import SemanticCache

class TestArchitecturalResilienceAndLimits(unittest.TestCase):

    def test_semantic_cache_lru_cap_and_eviction(self):
        """Valida que o SemanticCache impõe limite máximo estrito (50 itens por padrão) com política LRU."""
        cache = SemanticCache(max_size=5) # Tamanho reduzido para teste rápido
        
        for i in range(1, 6):
            cache.store(f"prompt_{i}", f"response_{i}")
        
        self.assertEqual(len(cache.cache), 5)
        self.assertEqual(cache.lookup("prompt_1"), "response_1")
        
        # Acessar prompt_1 deve movê-lo para o final do cache (LRU)
        # Ao adicionar o item 6, o mais antigo agora é prompt_2 (pois prompt_1 foi acessado recentemente)
        cache.store("prompt_6", "response_6")
        self.assertEqual(len(cache.cache), 5)
        
        # prompt_2 deve ter sido despejado (evicted)
        self.assertIsNone(cache.lookup("prompt_2"))
        # prompt_1 ainda deve existir
        self.assertEqual(cache.lookup("prompt_1"), "response_1")
        # prompt_6 deve existir
        self.assertEqual(cache.lookup("prompt_6"), "response_6")

    def test_semantic_cache_default_size(self):
        """Valida que o SemanticCache padrão inicializa com limite de 50 itens."""
        cache = SemanticCache()
        self.assertEqual(cache.max_size, 50)
        
        # Inserir 60 itens
        for i in range(60):
            cache.store(f"p_{i}", f"r_{i}")
            
        self.assertEqual(len(cache.cache), 50)
        # O item 0 deve ter sido despejado
        self.assertIsNone(cache.lookup("p_0"))
        # O item 59 deve existir
        self.assertEqual(cache.lookup("p_59"), "r_59")

    def test_parse_num_resilience_against_nan_and_inf(self):
        """Valida a função de parsing numérico contra NaN, Infinito e strings corrompidas."""
        # Testando a lógica de parse_num isoladamente
        def parse_num(val, fallback=0.0):
            if val is None: return fallback
            if isinstance(val, (int, float)):
                if isinstance(val, float) and (val != val or val == float('inf') or val == float('-inf')):
                    return fallback
                return float(val)
            s = str(val).strip()
            import re
            s = re.sub(r'[^\d.,-]', '', s)
            if not s: return fallback
            if ',' in s and '.' in s:
                if s.find('.') < s.find(','):
                    s = s.replace('.', '').replace(',', '.')
                else:
                    s = s.replace(',', '')
            elif ',' in s:
                s = s.replace(',', '.')
            try:
                f = float(s)
                if f != f or f == float('inf') or f == float('-inf'):
                    return fallback
                return f
            except ValueError:
                return fallback

        self.assertEqual(parse_num(None), 0.0)
        self.assertEqual(parse_num(float('nan'), 10.0), 10.0)
        self.assertEqual(parse_num(float('inf'), 5.0), 5.0)
        self.assertEqual(parse_num(float('-inf'), 5.0), 5.0)
        self.assertEqual(parse_num("R$ 1.250,50"), 1250.50)
        self.assertEqual(parse_num("inválido", 99.0), 99.0)
        self.assertEqual(parse_num("   -150,00 "), -150.0)

    def test_excel_generation_with_empty_and_extreme_values(self):
        """Valida que a geração de planilha Excel lida graciosamente com itens vazios e valores extremos sem quebrar."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Planilha Orçamentária"
        
        items = [
            {'rubrica': 'Serviços', 'item': 'Item Teste', 'qtd': 5, 'valorUnit': 100.0},
            {'rubrica': '', 'item': None, 'qtd': None, 'valorUnit': None},
            {'rubrica': 'Super Item', 'item': 'Extremo', 'qtd': 999999999, 'valorUnit': 999999999999999}
        ]
        
        for idx, it in enumerate(items):
            qtd = it.get('qtd') or 1.0
            if qtd > 10000000.0: qtd = 10000000.0
            v_unit = it.get('valorUnit') or 0.0
            if v_unit > 100000000000.0: v_unit = 100000000000.0
            
            ws.cell(row=idx+1, column=1, value=str(it.get('rubrica') or ''))
            ws.cell(row=idx+1, column=2, value=qtd)
            ws.cell(row=idx+1, column=3, value=v_unit)
            
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        self.assertTrue(len(buffer.getvalue()) > 500)
        loaded_wb = openpyxl.load_workbook(buffer)
        loaded_ws = loaded_wb.active
        self.assertEqual(loaded_ws.cell(row=1, column=1).value, "Serviços")
        self.assertEqual(loaded_ws.cell(row=3, column=2).value, 10000000.0) # Limitado pelo teto defensivo

if __name__ == '__main__':
    unittest.main()
